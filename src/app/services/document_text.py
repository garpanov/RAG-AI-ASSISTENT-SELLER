"""Text extraction from uploaded knowledge document files."""

from io import BytesIO
from pathlib import Path
from typing import ClassVar
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class UnsupportedDocumentFileError(ValueError):
    pass


class InvalidDocumentFileError(ValueError):
    pass


class EmptyDocumentFileError(ValueError):
    pass


class DocumentTextExtractor:
    """Extract text from supported in-memory files without persisting them."""

    _SUPPORTED_EXTENSIONS: ClassVar[set[str]] = {".pdf", ".docx", ".txt"}

    def extract(self, *, filename: str, data: bytes) -> str:
        extension = Path(filename).suffix.lower()
        if extension not in self._SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentFileError
        if not data:
            raise EmptyDocumentFileError

        try:
            if extension == ".pdf":
                text = self._extract_pdf(data)
            elif extension == ".docx":
                text = self._extract_docx(data)
            else:
                text = data.decode("utf-8-sig")
        except (
            BadZipFile,
            PdfReadError,
            PackageNotFoundError,
            UnicodeDecodeError,
            ValueError,
        ) as exc:
            raise InvalidDocumentFileError from exc

        text = text.strip()
        if not text:
            raise EmptyDocumentFileError
        return text

    @staticmethod
    def _extract_pdf(data: bytes) -> str:
        reader = PdfReader(BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    @staticmethod
    def _extract_docx(data: bytes) -> str:
        document = Document(BytesIO(data))
        parts = [
            paragraph.text for paragraph in document.paragraphs if paragraph.text
        ]
        parts.extend(
            " | ".join(cell.text for cell in row.cells if cell.text)
            for table in document.tables
            for row in table.rows
            if any(cell.text for cell in row.cells)
        )
        return "\n".join(parts)
