from io import BytesIO

import pytest
from docx import Document

from app.services.document_text import (
    DocumentTextExtractor,
    EmptyDocumentFileError,
    InvalidDocumentFileError,
    UnsupportedDocumentFileError,
)


def test_extracts_utf8_text() -> None:
    result = DocumentTextExtractor().extract(
        filename="policy.TXT",
        data=b"Return policy\nfor all orders",
    )

    assert result == "Return policy\nfor all orders"


def test_extracts_docx_paragraphs() -> None:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Return policy")
    document.add_paragraph("for all orders")
    document.save(buffer)

    result = DocumentTextExtractor().extract(
        filename="policy.docx",
        data=buffer.getvalue(),
    )

    assert result == "Return policy\nfor all orders"


def test_extracts_docx_tables() -> None:
    buffer = BytesIO()
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Period"
    table.cell(0, 1).text = "30 days"
    document.save(buffer)

    result = DocumentTextExtractor().extract(
        filename="policy.docx",
        data=buffer.getvalue(),
    )

    assert result == "Period | 30 days"


def test_rejects_unsupported_extension() -> None:
    with pytest.raises(UnsupportedDocumentFileError):
        DocumentTextExtractor().extract(filename="policy.csv", data=b"content")


def test_rejects_empty_text_file() -> None:
    with pytest.raises(EmptyDocumentFileError):
        DocumentTextExtractor().extract(filename="policy.txt", data=b" \n")


def test_rejects_invalid_pdf() -> None:
    with pytest.raises(InvalidDocumentFileError):
        DocumentTextExtractor().extract(filename="policy.pdf", data=b"not a pdf")


def test_rejects_invalid_docx() -> None:
    with pytest.raises(InvalidDocumentFileError):
        DocumentTextExtractor().extract(filename="policy.docx", data=b"not a docx")
